"""
Phase 8 - Editable XLSX report (openpyxl).
"""
import json
from io import BytesIO
from typing import Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from compliance_engine.reports.base import ENGINE_VERSION, _now_iso, _calibration_text
from compliance_engine.schema import ComplianceReport


def _style_header(ws, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(vertical="top")


def render_xlsx(
    report: ComplianceReport,
    out: Union[str, BytesIO, None] = None,
) -> Union[bytes, None]:
    """Render the editable XLSX workbook (Summary/Declarations/Violations/Font)."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Compliance status", report.compliance_status or ""])
    ws.append(["Engine version", ENGINE_VERSION])
    ws.append(["Generated (UTC)", _now_iso()])
    ws.append(["Calibration", _calibration_text(report)])
    stats = report.stats
    if stats:
        ws.append(["Total checks", stats.total_checks])
        ws.append(["Compliant", stats.compliant])
        ws.append(["Non-compliant", stats.non_compliant])
        ws.append(["Missing", stats.missing])
        ws.append(["Needs review", stats.needs_review])
        ws.append(["Not applicable", stats.not_applicable])
    _style_header(ws, 2)

    ws = wb.create_sheet("Declarations")
    ws.append(["Field", "Value", "Source zone", "Confidence", "Method"])
    for d in report.declarations:
        ws.append([d.field_name, d.value, d.source_zone, d.confidence, d.method])
    _style_header(ws, 5)

    ws = wb.create_sheet("Violations")
    ws.append(["Rule ID", "Status", "Severity", "Reason", "Rule reference",
               "Evidence"])
    for v in report.violations:
        ws.append([v.rule_id, v.status, v.severity, v.reason,
                   v.rule_reference,
                   json.dumps(v.evidence, default=str)[:400]])
    _style_header(ws, 6)

    ws = wb.create_sheet("FontMetrics")
    ws.append(["Zone type", "Char height px", "Char height mm", "Contrast",
               "Calibrated"])
    for m in report.font_metrics:
        ws.append([m.zone_type, m.char_height_px_median,
                   m.char_height_mm_median, m.contrast_ratio, m.calibrated])
    _style_header(ws, 5)

    for sheet in wb.worksheets:
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 26

    if out is None:
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
    if isinstance(out, str):
        wb.save(out)
        return None
    wb.save(out)
    return None