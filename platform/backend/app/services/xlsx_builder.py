"""Editable XLSX compliance workbook (openpyxl).

Sheets: Summary / Product / Declarations / Violations / FontMetrics / Evidence.
Every sheet has a frozen header row and auto-filters so officers can filter,
clean and re-use the workbook as editable evidence.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.inspection import Inspection
from app.services.report_meta import report_no

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_WRAP = Alignment(vertical="top", wrap_text=True)


def _style(ws, widths):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_xlsx(inspection: Inspection) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    stats = inspection.stats_json or {}
    ws.append(["Report", "LMPC Compliance Report"])
    ws.append(["Report number", _report_no(inspection)])
    ws.append(["Compliance status", inspection.compliance_status])
    ws.append(["Generated (UTC)", inspection.updated_at.isoformat() if inspection.updated_at else ""])
    ws.append(["Image", inspection.image_name])
    ws.append(["Engine scan id", inspection.engine_scan_id])
    ws.append(["Total checks", stats.get("total_checks", "")])
    ws.append(["Compliant", stats.get("compliant", "")])
    ws.append(["Non-compliant", stats.get("non_compliant", "")])
    ws.append(["Missing", stats.get("missing", "")])
    ws.append(["Needs review", stats.get("needs_review", "")])
    ws.append(["Not applicable", stats.get("not_applicable", "")])
    _style(ws, [24, 60])

    ws = wb.create_sheet("Product")
    p = inspection.product
    ws.append(["Field", "Value"])
    rows = [("Generic name", p.generic_name if p else ""), ("Brand", p.brand if p else ""),
            ("Category", p.category if p else ""), ("Manufacturer", p.manufacturer if p else ""),
            ("Packer", p.packer if p else ""), ("Importer", p.importer if p else ""),
            ("Net quantity", p.net_quantity_text if p else ""), ("MRP", p.mrp if p else ""),
            ("Batch no", p.batch_no if p else "")]
    for r in rows:
        ws.append(list(r))
    _style(ws, [24, 60])

    ws = wb.create_sheet("Declarations")
    ws.append(["Field", "Value", "Raw text", "Source zone", "Confidence", "Method", "Qualifiers"])
    for d in inspection.declarations:
        ws.append([d.field_name, d.value, d.raw_text, d.source_zone, d.confidence,
                   d.method, " | ".join(str(q) for q in (d.qualifiers or []))])
    _style(ws, [36, 40, 40, 22, 12, 14, 30])

    ws = wb.create_sheet("Violations")
    ws.append(["Rule ID", "Rule reference", "Field", "Status", "Severity", "Reason", "Evidence"])
    for v in inspection.violations:
        ws.append([v.rule_id, v.rule_reference, v.field_name, v.status, v.severity,
                   v.reason, _compact(v.evidence)])
    _style(ws, [16, 34, 26, 18, 12, 70, 50])

    ws = wb.create_sheet("FontMetrics")
    ws.append(["Zone", "Char height (px)", "Char height (mm)", "Contrast", "Calibrated"])
    for m in inspection.font_metrics:
        ws.append([m.zone_type, m.char_height_px_median, m.char_height_mm_median,
                   m.contrast_ratio, m.calibrated])
    _style(ws, [34, 18, 18, 14, 12])

    ws = wb.create_sheet("Evidence")
    ws.append(["File", "MIME", "Size (bytes)", "Uploaded by", "Uploaded at"])
    for e in inspection.evidence:
        ws.append([e.filename, e.mime_type, e.size_bytes, e.uploaded_by, e.created_at])
    _style(ws, [40, 24, 16, 14, 26])

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = _WRAP

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _report_no(inspection: Inspection) -> str:
    ts = (inspection.created_at or "").strftime("%Y%m%d") if hasattr(inspection.created_at, "strftime") else "00000000"
    return f"LMPC-{ts}-{inspection.token[:6].upper()}"


def _compact(value) -> str:
    import json

    try:
        return json.dumps(value, default=str)[:500]
    except Exception:
        return str(value)[:500]